"""Offline integration: real template tools, writing loop, checkpoints and DOCX export."""
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path

import pytest
from langchain_core.messages import AIMessage

from agent.workflow import PaperWorkflow
from agent.writing import WritingAgent, WritingExecutionError
from agent.research import assert_research_lock, freeze_research
from agent.assembly.engine import NS, package, parse, text
from tests.support.paper_fixture import fixture


def test_visual_consistency_gate_attaches_render_and_rejects_crop(tmp_path):
    import base64
    from PIL import Image
    from agent.research import review_consistency
    path = tmp_path / "diagram.png"
    Image.new("RGB", (32, 32), "white").save(path)
    class Reviewer:
        def bind(self, **kwargs):
            return self
        def invoke(self, messages):
            parts = messages[1].content
            assert json.loads(parts[0]["text"])["content"] == {"sections": []}
            url = parts[2]["image_url"]["url"]
            assert url.startswith("data:image/png;base64,")
            assert base64.b64decode(url.split(",", 1)[1]) == path.read_bytes()
            return AIMessage(content='{"issues": ["fig_1 节点被裁切"]}')
    state = {"topic": "校园平台"}
    freeze_research(state)
    with pytest.raises(ValueError, match="fig_1 节点被裁切"):
        review_consistency(Reviewer(), state, {"sections": []},
                           images=[{"asset_id": "fig_1", "path": str(path)}])


@pytest.mark.parametrize("confirmed", [[], ["订单状态值与冻结方案矛盾"]])
def test_review_rechecks_self_withdrawn_findings(confirmed):
    from agent.research import review_consistency
    class Reviewer:
        calls = 0
        def bind(self, **kwargs):
            return self
        def invoke(self, messages):
            self.calls += 1
            issues = ["订单状态值与冻结方案矛盾", "字段与方案一致，无冲突"] if self.calls == 1 else confirmed
            return AIMessage(content=json.dumps({"issues": issues}))
    state, model = {"topic": "校园平台"}, Reviewer()
    freeze_research(state)
    if confirmed:
        with pytest.raises(ValueError, match="订单状态值与冻结方案矛盾"):
            review_consistency(model, state, {})
    else:
        assert review_consistency(model, state, {})["passed"]
    assert model.calls == 2


def test_review_receives_approved_design_details_without_mutating_research():
    from agent.research import review_consistency
    record = {"section_id": "4.2.3", "summary": "拟议退款状态机",
              "key_facts": ["退款成功后由退款中转为已取消"]}
    state = {"topic": "校园平台", "writing_records": [record]}
    freeze_research(state)
    before = deepcopy(state["research_lock"])
    class Reviewer:
        def bind(self, **kwargs):
            return self
        def invoke(self, messages):
            context = json.loads(messages[1].content)
            assert context["prior_design"] == [record]
            assert "不是实测证据" in messages[0].content
            return AIMessage(content='{"issues": []}')
    assert review_consistency(Reviewer(), state, {"test_plan": "验证退款终态"})["passed"]
    assert state["research_lock"] == before


@pytest.mark.parametrize("accept_repair", [True, False])
def test_final_image_repair_is_reviewed_before_commit_and_resumes(tmp_path, accept_repair):
    from PIL import Image
    from agent.writing.asset_review import FinalAssetReviewAgent
    from agent.writing.storage import input_fingerprint
    from tests.support.paper_fixture import commit_fixture
    state, _ = fixture(tmp_path)
    assets = []
    (tmp_path / "writing/assets").mkdir(parents=True)
    for name in ("writing/assets/fig_a", "writing/assets/fig_b"):
        (tmp_path / (name + ".py")).write_text("# old")
        Image.new("RGB", (640, 320), "white").save(tmp_path / (name + ".png"))
        assets.append({"asset_id": Path(name).name, "description": "预约流程", "kind": "diagram",
                       "source_path": name + ".py", "path": name + ".png"})
    commit_fixture(state, assets)
    freeze_research(state)
    path = tmp_path / "writing/manifest.json"
    path.write_text(json.dumps({"version": 1, "input_fingerprint": input_fingerprint(state),
                                "records": state["writing_records"]}))
    before = path.read_bytes()
    class Reviewer:
        calls = 0
        def bind(self, **kwargs):
            return self
        def invoke(self, messages):
            self.calls += 1
            context = json.loads(messages[1].content[0]["text"])
            issues = [] if accept_repair and context["content"]["source"] == "# fixed" else ["文字被裁切"]
            return AIMessage(content=json.dumps({"issues": issues}))
    class Coder:
        def run_asset(self, task, workspace):
            assert task["previous_source_path"].endswith(".py")
            name = "writing/assets/" + task["asset_id"] + "_fixed"
            (tmp_path / (name + ".py")).write_text("# fixed")
            Image.new("RGB", (640, 320), "white").save(tmp_path / (name + ".png"))
            return {"asset_id": task["asset_id"], "description": "预约流程", "kind": "diagram",
                    "source_path": name + ".py", "path": name + ".png"}
    model = Reviewer()
    agent = FinalAssetReviewAgent(model, coder=Coder(), max_repairs=1)
    if not accept_repair:
        with pytest.raises(ValueError, match="成稿图表复查未通过"):
            agent.invoke(state)
        assert path.read_bytes() == before
    else:
        agent.invoke(state)
        assert all(a["path"].endswith("_fixed.png") for a in state["writing_records"][0]["assets"])
        assert state["asset_review_report"]["replacements"] == 2
        assert model.calls == 4
        agent.invoke(state)
        assert model.calls == 4
        assert (tmp_path / "writing/assets/fig_a.py").read_text() == "# old"


class Model:
    def __init__(self, source):
        self.source = source
        self.writes = 0
        self.fail_after = None
        self.reject = False
        self.calls = []

    def bind(self, **kwargs):
        return self

    def bind_tools(self, tools, **kwargs):
        return self

    def invoke(self, messages):
        context = json.loads(messages[1].content)
        self.calls.append(context)
        if "template_inventory" in context:
            ps = parse(package(Path(self.source["new_template_path"]))["word/document.xml"]).find("w:body", NS).findall("w:p", NS)
            # One custom style per body sample proves replace_all through the real script.
            definitions = [{"id": f"PaperSample{i}", "name": f"模板样式{i}",
                            "sample": {"paragraph": i}} for i in range(len(ps))]
            value = {"plan": {"edits": [], "style_rebuild": {"mode": "replace_all",
                      "styles": definitions, "default_style": "PaperSample0",
                      "assignments": [{"style": d["id"], "paragraphs": [i]} for i, d in enumerate(definitions)]}},
                     "assembly_contract": self.source["assembly_contract"]}
        elif "user_facts" in context:
            value = {"research_object": "预约系统", "objectives": ["避免时段冲突"],
                     "research_methods": ["需求分析", "用例设计"], "entities": [{"name": "预约"}],
                     "technical_design": {"language": "Python"}, "evaluation_plan": ["拟验证时段冲突"],
                     "limitations": ["未进行实测"]}
        elif "frozen_research" in context:
            value = {"issues": ["研究方法不一致"] if self.reject and "sections" in context["content"] else []}
        elif "current_tasks" in context:
            if self.fail_after is not None and self.writes >= self.fail_after:
                raise RuntimeError("模拟模型断线")
            self.writes += 1
            value = {"sections": [{"section_id": t["section_id"], "title": t["title"],
                     "summary": "预约方案与验证计划", "key_facts": ["Python"], "assets": [],
                     "citations": ["ref1"], "content": "本研究拟设计预约系统，使用Python描述服务端逻辑。"
                     "方案要求校验预约时段，学生提交请求后应检查是否存在冲突，管理员负责维护规则。"
                     "后续拟通过预约用例验证设计，目前尚无实测结果。相关主题见[[REF_CITE:ref1]]。"}
                     for t in context["current_tasks"]]}
        elif "research" in context:
            value = {"abstract": "本研究提出预约系统设计方案，拟验证时段冲突，尚未实测。", "keywords": ["预约", "系统设计"]}
        else:
            raise AssertionError(context)
        return AIMessage(content=json.dumps(value, ensure_ascii=False))


class Literature:
    def __init__(self):
        self.calls = 0

    def invoke(self, state):
        self.calls += 1
        state["bib_pool"] = [{"key": "ref1", "title": "预约系统研究（测试夹具）", "authors": "测试作者",
            "journal": "测试期刊", "year": 2024, "language": "zh", "language_basis": "title_script",
            "doi": None, "url": "https://example.org/paper", "source": "crossref", "source_id": "fixture",
            "source_url": "https://example.org/source", "query": "预约系统", "retrieved_at": "2026-01-01"}]


def setup(tmp_path):
    source, _ = fixture(tmp_path / "source")
    state = {"topic": "预约系统设计", "raw_doc_path": source["new_template_path"],
             "workspace_dir": str(tmp_path / "task"), "single_source_of_truth": {"tech_stack": "Python"},
             "outline_plan": [{"section_id": "1", "title": "第一章 研究设计", "target_words": 0,
                               "status": "pending", "planned_assets": []},
                              {"section_id": "1.1", "title": "设计方案", "target_words": 0,
                               "status": "pending", "planned_assets": []}] + [
                 {"section_id": f"1.1.{i}", "title": f"设计细节{i}", "target_words": 100,
                  "status": "pending", "planned_assets": []} for i in range(1, 5)]}
    model = Model(source)
    literature = Literature()
    workflow = PaperWorkflow(llm=model, literature=literature,
                             writer=WritingAgent(llm=model, skills=[], max_repairs=0))
    return state, model, literature, workflow


def test_full_workflow_resume_and_template_preservation(tmp_path):
    state, model, literature, workflow = setup(tmp_path)
    original = deepcopy(state)
    source_hash = sha256(Path(state["raw_doc_path"]).read_bytes()).hexdigest()
    model.fail_after = 1
    with pytest.raises(WritingExecutionError, match="断线"):
        workflow.invoke(state)
    checkpoint = Path(state["workspace_dir"]) / "workflow/state.json"
    failed = json.loads(checkpoint.read_text())["state"]
    assert failed["workflow_report"]["current_stage"] == "writing"
    assert failed["workflow_report"]["status"] == "failed"
    assert "writing" not in failed["workflow_report"]["completed_stages"]
    assert len(json.loads((checkpoint.parents[1] / "writing/manifest.json").read_text())["records"]) == 3
    model.fail_after = None
    result = workflow.invoke(state)
    assert state == original
    assert literature.calls == 1
    assert model.writes == 2
    assert len(result["writing_records"]) == 4
    assert result["workflow_report"]["status"] == "completed"
    assert result["assembly_report"]["deterministic_replay_identical"]
    assert result["assembly_report"]["protected_parts_identical"]
    assert source_hash == sha256(Path(state["raw_doc_path"]).read_bytes()).hexdigest()
    assert_research_lock(result)
    clean = package(Path(result["new_template_path"]))
    output = package(Path(result["final_doc_path"]))
    assert clean["word/styles.xml"] == output["word/styles.xml"]
    assert all(s.get("{" + NS["w"] + "}customStyle") == "1" for s in parse(clean["word/styles.xml"]).findall("w:style", NS))
    rendered = text(parse(output["word/document.xml"]))
    assert "{{" not in rendered
    assert "1.1.1 设计细节1" in rendered
    calls = len(model.calls)
    again = workflow.invoke(state)
    assert len(model.calls) == calls
    assert again["assembly_report"]["output_sha256"] == result["assembly_report"]["output_sha256"]
    with pytest.raises(ValueError, match="不同输入"):
        workflow.invoke({**state, "topic": "换一个题目"})
    content = Path(result["workspace_dir"]) / result["writing_records"][0]["content_path"]
    content.write_text("被篡改")
    with pytest.raises(ValueError, match="被修改"):
        workflow.invoke(state)


def test_conflicting_batch_never_commits(tmp_path):
    state, model, _, workflow = setup(tmp_path)
    model.reject = True
    with pytest.raises(WritingExecutionError, match="研究方法不一致"):
        workflow.invoke(state)
    root = Path(state["workspace_dir"])
    assert not (root / "writing/manifest.json").exists()
    assert not (root / "assembly/final.docx").exists()
    assert "研究方法不一致" in (root / "writing/last-batch-error.json").read_text()
    model.reject = False
    workflow.invoke(state)
    assert any(c.get("previous_validation_errors") for c in model.calls if "current_tasks" in c)


def test_frozen_entity_and_method_changes_are_rejected():
    state = {"topic": "原始选题", "single_source_of_truth": {"entities": ["用户"], "methods": ["用例测试"]}}
    freeze_research(state)
    state["single_source_of_truth"]["methods"].append("编造实验")
    with pytest.raises(ValueError, match="冻结后"):
        assert_research_lock(state)


def test_empty_literature_stops_before_writing(tmp_path):
    state, model, _, workflow = setup(tmp_path)
    class Empty:
        def invoke(self, state):
            state["bib_pool"] = []
    workflow.literature = Empty()
    with pytest.raises(ValueError, match="没有取得可核验文献"):
        workflow.invoke(state)
    assert model.writes == 0


def test_two_row_grid_template_reuses_data_row(tmp_path):
    from agent.assembly import AssemblyAgent
    from tests.support.paper_fixture import commit_fixture
    from docx import Document
    state, contract = fixture(tmp_path)
    template = Path(state["new_template_path"])
    doc = Document(template)
    doc.tables[0]._tbl.remove(doc.tables[0].rows[-1]._tr)
    doc.save(template)
    contract["template_sha256"] = sha256(template.read_bytes()).hexdigest()
    folder = tmp_path / "writing/assets/table"
    folder.mkdir(parents=True)
    data = folder / "table.json"
    data.write_text(json.dumps({"columns": ["对象", "说明"], "rows": [["A", "一"], ["B", "二"], ["C", "三"]]}))
    commit_fixture(state, [{"asset_id": "tbl_1", "kind": "table", "description": "对象表",
        "source_path": str(data.relative_to(tmp_path)), "path": str(data.relative_to(tmp_path))}])
    report = AssemblyAgent(state, contract, "研究方案", ["方案"], str(tmp_path / "assembly/final.docx")).invoke()
    assert report["deterministic_replay_identical"]
    result = Document(report["path"])
    assert len(result.tables[0].rows) == 4
    assert result.tables[0].cell(3, 1).text == "三"


def test_bilingual_template_metadata_is_not_filled_with_chinese(tmp_path):
    from agent.assembly import AssemblyAgent
    from tests.support.paper_fixture import commit_fixture
    from docx import Document
    state, contract = fixture(tmp_path)
    template = Path(state["new_template_path"])
    doc = Document(template)
    for marker, key in [("{{EnglishTitle}}", "topic_en"), ("{{EnglishAbstract}}", "abstract_en"),
                        ("{{EnglishKeywords}}", "keywords_en")]:
        doc.add_paragraph(marker)
        contract["fields"][marker] = key
    doc.save(template)
    contract["template_sha256"] = sha256(template.read_bytes()).hexdigest()
    commit_fixture(state)
    state.update(topic_en="Appointment System Design", abstract_en="This study proposes an appointment system.",
                 keywords_en=["appointment", "design"])
    report = AssemblyAgent(state, contract, "本文提出预约系统方案。", ["预约", "设计"], str(tmp_path / "assembly/final.docx")).invoke()
    content = text(parse(package(Path(report["path"]))["word/document.xml"]))
    assert "Appointment System Design" in content
    assert "This study proposes an appointment system." in content
    assert "appointment; design" in content
    assert "{{" not in content


def test_hierarchy_normalization_preserves_all_leaf_assets_and_titles():
    from agent.output_parsers.outline_normalization import normalize_workflow_outline
    def node(sid, title, words=0, assets=None):
        return {"section_id": sid, "title": title, "target_words": words,
                "status": "pending", "planned_assets": assets or []}
    source = [node("1", "绪论"), node("1.1", "背景", 1000), node("2", "设计"), node("2.1", "数据库"),
              node("2.1.1", "概念", 10000, ["fig_er"]), node("2.1.2", "逻辑"),
              node("2.1.2.1", "用户表", 10000, ["tbl_user"]), node("2.1.2.2", "订单表", 10000, ["tbl_order"])]
    raw, report = normalize_workflow_outline(json.dumps({"outline_plan": source}))
    result = json.loads(raw)["outline_plan"]
    assert result[-2]["title"] == "逻辑：用户表"
    assert result[-1]["title"] == "逻辑：订单表"
    assert [n["section_id"] for n in result[-2:]] == ["2.1.2", "2.1.3"]
    assert [a for n in result for a in n["planned_assets"]] == ["fig_er", "tbl_user", "tbl_order"]
    assert max(n["section_id"].count(".") for n in result) == 2
    assert 11000 <= report["target_words"] <= 13000
    assert source[-1]["section_id"] == "2.1.2.2"


def test_reader_schema_only_advertises_committed_sections(tmp_path):
    from agent.writing.tools import create_writing_tools
    tools = {t.name: t for t in create_writing_tools(str(tmp_path), records=[
        {"section_id": "1.1", "content": "已提交的正文"}])}
    assert tools["read_section"].args["section_id"]["enum"] == ["1.1"]
    assert tools["read_section"].invoke({"section_id": "1.1"}) == "已提交的正文"
    with pytest.raises(PermissionError):
        tools["read_section"].invoke({"section_id": "2.1"})


def test_recent_committed_sections_remain_readable_and_reads_are_bounded(tmp_path):
    state, base, literature, _ = setup(tmp_path)
    leaf = state["outline_plan"][-1]
    state["outline_plan"].extend({**leaf, "section_id": f"1.1.{i}", "title": f"设计细节{i}"}
                                 for i in range(5, 8))

    class Reader(Model):
        def __init__(self):
            super().__init__(base.source)
            self.reads = 0

        def bind_tools(self, tools, **kwargs):
            owner = self
            reader = next(t for t in tools if t.name == "read_section")
            assert "1.1.6" in reader.args["section_id"]["enum"]
            class Bound:
                def invoke(self, messages):
                    owner.reads += 1
                    assert reader.invoke({"section_id": "1.1.6"})
                    return AIMessage(content="", tool_calls=[{"name": "read_section",
                        "args": {"section_id": "1.1.6"}, "id": f"read{owner.reads}"}])
            return Bound()

    model = Reader()
    result = PaperWorkflow(llm=model, literature=literature,
        writer=WritingAgent(llm=model, skills=[])).invoke(state)
    assert len(result["writing_records"]) == 7
    assert model.reads == 2


def test_asset_corrections_regenerate_instead_of_reusing_rejected_data(tmp_path):
    from agent.writing.tools import create_writing_tools
    class Coder:
        def __init__(self):
            self.instructions = []
        def run_asset(self, task, workspace):
            self.instructions.append(task["instructions"])
            return {"description": task["instructions"]}
    coder = Coder()
    issued = {}
    tools = {t.name: t for t in create_writing_tools(str(tmp_path),
        tasks=[{"section_id": "2.1", "planned_assets": ["fig_design: 方案图"]}], coder=coder, issued=issued)}
    call = {"section_id": "2.1", "asset_id": "fig_design", "instructions": "初始方案"}
    tools["code_asset"].invoke(call)
    tools["code_asset"].invoke(call)
    tools["code_asset"].invoke({**call, "instructions": "删除没有事实依据的新增实体"})
    assert coder.instructions == ["初始方案", "删除没有事实依据的新增实体"]
    assert issued[("2.1", "fig_design")]["description"] == coder.instructions[-1]


def test_coder_reuses_verified_runtime_profile_and_publishes_real_assets(tmp_path):
    from agent.coding import CodingAgent
    class CodingModel:
        def bind_tools(self, tools):
            return self
        def invoke(self, messages):
            context = json.loads(messages[1].content)
            assert context["runtime_environment"]["modules"]["PIL"]
            if context["asset_id"] == "tbl_2":
                assert context["previous_source"] == "previous_source.py"
                assert "previous_source_path" not in context
            if len(messages) == 2:
                return AIMessage(content="", tool_calls=[{"name": "bash", "id": "generate", "args": {
                    "command": "cat > table.py <<'PY'\n# version " + context["asset_id"] + "\nimport json\nfrom pathlib import Path\nPath('table.json').write_text(json.dumps({'columns':['字段','类型','说明'],'rows':[['id','Long','主键']]}, ensure_ascii=False))\nPY\npython table.py"}}])
            assert json.loads(next(m.content for m in reversed(messages) if m.type == "tool"))["exit_code"] == 0
            assert "产物提交阶段" in messages[-1].content
            return AIMessage(content=json.dumps({"source_path": "table.py", "path": "table.json",
                                                 "kind": "table", "description": "字段方案"}))
    coder = CodingAgent(llm=CodingModel())
    previous = None
    for number in (1, 2):
        asset = coder.run_asset({"section_id": "2.1", "asset_id": f"tbl_{number}",
                                "previous_source_path": previous,
                                "render_constraints": {"table_columns": [3]}}, str(tmp_path))
        assert (tmp_path / asset["path"]).is_file()
        assert (tmp_path / asset["source_path"]).is_file()
        previous = asset["source_path"]
    assert next((tmp_path / "subagents/coding").rglob("previous_source.py")).read_text().startswith("# version tbl_1")
    traces = [json.loads(p.read_text()) for p in (tmp_path / "audit/coding").glob("*.json")]
    assert sum("runtime_probe" in event for trace in traces for event in trace) == 1


def test_relationship_diagram_uses_executable_dot_source(tmp_path):
    import shutil
    if not shutil.which("dot"):
        pytest.skip("optional Graphviz dependency is absent")
    from agent.coding import CodingAgent
    class Model:
        calls = 0
        def bind_tools(self, tools):
            return self
        def invoke(self, messages):
            self.calls += 1
            assert json.loads(messages[1].content)["layout_engine"] == "graphviz"
            if self.calls == 1:
                return AIMessage(content="", tool_calls=[{"name": "bash", "id": "draw", "args": {
                    "command": "printf 'digraph G { a -> b; }' > graph.dot\nprintf '# not the graph source' > wrong.py\ndot -Tpng graph.dot -o graph.png"}}])
            if self.calls == 3:
                assert "DOT" in messages[-1].content
            return AIMessage(content=json.dumps({"source_path": "wrong.py" if self.calls == 2 else "graph.dot",
                "path": "graph.png", "kind": "diagram", "description": "实体关系图"}))
    model = Model()
    asset = CodingAgent(llm=model).run_asset({"section_id": "2.1", "asset_id": "fig_1", "requirement": "ER图"}, str(tmp_path))
    assert model.calls == 3
    assert asset["source_path"].endswith(".dot")


def test_coding_pipeline_does_not_hide_renderer_failure(tmp_path):
    from agent.security.sandbox import CodingSandbox
    result = json.loads(CodingSandbox(tmp_path).run("python -c 'raise RuntimeError(\"render failed\")' 2>&1 | tail -1"))
    assert result["exit_code"] != 0
    assert "render failed" in result["stdout"]


def test_figure_fits_template_width_without_distorting_content(tmp_path):
    from PIL import Image
    from docx import Document
    from agent.assembly import AssemblyAgent
    from agent.assembly.engine import asset_constraints
    from tests.support.paper_fixture import commit_fixture
    state, contract = fixture(tmp_path)
    folder = tmp_path / "writing/assets"
    folder.mkdir(parents=True)
    Image.new("RGB", (600, 800), "white").save(folder / "diagram.png")
    (folder / "source.py").write_text("# 测试图表源码")
    commit_fixture(state, [{"asset_id": "fig_1", "kind": "diagram", "description": "竖向流程图",
        "path": "writing/assets/diagram.png", "source_path": "writing/assets/source.py"}])
    constraints = asset_constraints(state)
    assert "image_aspect_ratio" not in constraints
    output = tmp_path / "assembly/final.docx"
    AssemblyAgent(state, contract, "摘要", ["设计"], str(output)).invoke()
    shape = Document(output).inline_shapes[0]
    assert shape.width / shape.height == pytest.approx(600 / 800, rel=0.001)
    assert shape.width <= constraints["image_max_width_emu"]
    assert shape.height <= constraints["image_max_height_emu"]
    original = package(Path(state["new_template_path"]))
    final = package(output)
    assert original["word/styles.xml"] == final["word/styles.xml"]
    contract["image_sizing"] = "fixed"
    with pytest.raises(ValueError, match="宽高比"):
        AssemblyAgent(state, contract, "摘要", ["设计"], str(output)).invoke()


def test_resume_reuses_prepared_assets_after_provider_failure(tmp_path):
    state, base, literature, _ = setup(tmp_path)
    state["outline_plan"][2]["planned_assets"] = ["tbl_roles 权限方案"]
    class Coder:
        calls = 0
        def run_asset(self, task, workspace):
            self.calls += 1
            path = Path(workspace) / "writing/assets/roles.json"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps({"columns": ["角色", "职责"], "rows": [["学生", "预约"]]}))
            return {"asset_id": "tbl_roles", "kind": "table", "description": "权限方案",
                    "source_path": "writing/assets/roles.json", "path": "writing/assets/roles.json"}
    class AssetModel(Model):
        fail = True
        tool_names = ()
        def bind_tools(self, tools, **kwargs):
            self.tool_names = [t.name for t in tools]
            return self
        def invoke(self, messages):
            context = json.loads(messages[1].content)
            if "current_tasks" not in context or not context["current_tasks"][0]["planned_assets"]:
                return super().invoke(messages)
            cached = context.get("prepared_assets", [])
            if not cached and len(messages) == 2:
                return AIMessage(content="", tool_calls=[{"name": "code_asset", "id": "asset",
                    "args": {"section_id": "1.1.1", "asset_id": "tbl_roles", "instructions": "按角色方案制表"}}])
            if self.fail:
                assert any("本批规划资产已全部生成" in str(m.content) for m in messages)
                raise RuntimeError("模拟图表完成后模型额度耗尽")
            assert cached
            assert any("本批规划资产已全部生成" in str(m.content) for m in messages)
            response = super().invoke(messages)
            value = json.loads(response.content)
            value["sections"][0]["assets"] = [cached[0]["asset"]]
            value["sections"][0]["content"] += "如[[REF_TABLE:tbl_roles]]所示。"
            return AIMessage(content=json.dumps(value, ensure_ascii=False))
    coder, model = Coder(), AssetModel(base.source)
    workflow = PaperWorkflow(llm=model, literature=literature,
        writer=WritingAgent(llm=model, skills=[], coder=coder))
    with pytest.raises(WritingExecutionError, match="额度耗尽"):
        workflow.invoke(state)
    root = Path(state["workspace_dir"])
    assert (root / "writing/prepared-assets.json").exists()
    assert not (root / "writing/manifest.json").exists()
    model.fail = False
    result = workflow.invoke(state)
    assert coder.calls == 1
    assert len(result["writing_records"]) == 4
    assert result["assembly_report"]["assets"] == 1
    from agent.writing.prepared_assets import load_prepared
    from agent.assembly.engine import asset_constraints
    (root / "writing/assets/roles.json").write_text(json.dumps({"columns": ["角色"], "rows": [["修改"]]}))
    with pytest.raises(ValueError, match="缓存被修改"):
        load_prepared(root, result, WritingAgent._tasks(result), asset_constraints(result))


def test_coding_provider_quota_stops_batch_without_repeated_delegation(tmp_path):
    state, base, literature, _ = setup(tmp_path)
    state["outline_plan"][2]["planned_assets"] = ["tbl_roles 权限方案"]
    class QuotaError(Exception):
        status_code = 402
    class Coder:
        calls = 0
        def run_asset(self, task, workspace):
            self.calls += 1
            raise QuotaError("模型余额不足")
    class AssetModel(Model):
        def invoke(self, messages):
            context = json.loads(messages[1].content)
            if "current_tasks" in context:
                return AIMessage(content="", tool_calls=[{"name": "code_asset", "id": "asset",
                    "args": {"section_id": "1.1.1", "asset_id": "tbl_roles", "instructions": "生成表格"}}])
            return super().invoke(messages)
    coder, model = Coder(), AssetModel(base.source)
    workflow = PaperWorkflow(llm=model, literature=literature,
                            writer=WritingAgent(llm=model, skills=[], coder=coder))
    with pytest.raises(WritingExecutionError, match="模型余额不足"):
        workflow.invoke(state)
    assert coder.calls == 1
    assert not (Path(state["workspace_dir"]) / "writing/manifest.json").exists()


def test_coder_rejects_binary_image_as_source_and_requests_repair(tmp_path):
    from agent.coding import CodingAgent
    class Model:
        calls = 0
        def bind_tools(self, tools):
            return self
        def invoke(self, messages):
            self.calls += 1
            if self.calls == 1:
                return AIMessage(content="", tool_calls=[{"name": "bash", "id": "draw", "args": {
                    "command": "cat > source.py <<'PY'\nfrom PIL import Image\nImage.new('RGB',(100,50),'white').save('output.png')\nPY\npython source.py"}}])
            return AIMessage(content=json.dumps({"source_path": "output.png" if self.calls == 2 else "source.py",
                "path": "output.png", "kind": "diagram", "description": "测试图"}))
    model = Model()
    asset = CodingAgent(llm=model).run_asset({"section_id": "2.1", "asset_id": "fig_1"}, str(tmp_path))
    assert model.calls == 3
    assert asset["source_path"].endswith("source.py")
    trace = json.loads(next((tmp_path / "audit/coding").glob("*.json")).read_text())
    assert any("UTF-8" in e.get("validation_error", "") for e in trace)


def test_export_requests_toc_refresh_without_changing_other_settings(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from lxml import etree
    from agent.assembly import AssemblyAgent
    from tests.support.paper_fixture import commit_fixture
    state, contract = fixture(tmp_path)
    template = Path(state["new_template_path"])
    doc = Document(template)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), ' TOC \\o "1-3" ')
    doc.add_paragraph()._p.append(field)
    doc.save(template)
    contract["template_sha256"] = sha256(template.read_bytes()).hexdigest()
    commit_fixture(state)
    output = tmp_path / "assembly/final.docx"
    report = AssemblyAgent(state, contract, "摘要", ["设计"], str(output)).invoke()
    before, after = package(template), package(output)
    assert report["update_fields_on_open"]
    assert before["word/styles.xml"] == after["word/styles.xml"]
    old, new = parse(before["word/settings.xml"]), parse(after["word/settings.xml"])
    update = new.find("w:updateFields", NS)
    assert update.get(qn("w:val")) == "true"
    new.remove(update)
    assert etree.tostring(old) == etree.tostring(new)
    assert parse(after["word/document.xml"]).xpath(".//w:fldSimple[@w:dirty='true']", namespaces=NS)
