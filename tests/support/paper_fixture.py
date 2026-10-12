"""构造带非默认格式、多分节、页眉页脚、图表的批准模板。"""
from hashlib import sha256
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_SECTION_START
from PIL import Image


def fixture(root):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    image = root / "template-image.png"
    Image.new("RGB", (640, 320), "white").save(image)
    doc = Document()
    from docx.oxml.ns import qn
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"].font.size = Pt(11)
    doc.sections[0].left_margin = Cm(3.17)
    doc.sections[0].header.paragraphs[0].text = "学校模板 {{论文题目}}"
    doc.sections[0].footer.paragraphs[0].text = "保留页脚格式"
    doc.add_paragraph("{{论文题目}}")
    doc.add_paragraph("{{摘要}}")
    doc.add_paragraph("{{关键词}}")
    doc.add_paragraph("学校声明：不得删除")
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph("{{正文}}")
    mapping = {}
    for name, style in (("heading1", "Heading 1"), ("heading2", "Heading 2"),
                        ("heading3", "Heading 3"), ("body", "Normal"),
                        ("figure_caption", "Caption"), ("table_caption", "Caption"),
                        ("bibliography_heading", "Heading 1"), ("bibliography", "Normal")):
        marker = "{{原型:"+name+"}}"
        p = doc.add_paragraph(marker, style)
        p.paragraph_format.space_after = Pt(7)
        p.runs[0].font.name = "宋体"
        mapping[name] = marker
    p = doc.add_paragraph("{{原型:figure}}")
    p.add_run().add_picture(str(image), width=Cm(12), height=Cm(6))
    mapping["figure"] = "{{原型:figure}}"
    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Cm(4+i)
            cell.text = "原型单元格"
            cell.paragraphs[0].runs[0].font.size = Pt(9)
    template = root / "approved-template.docx"
    doc.save(template)
    contract = {"template_sha256": sha256(template.read_bytes()).hexdigest(), "body_marker": "{{正文}}",
                "paragraphs": mapping, "tables": {"2": 0},
                "fields": {"{{论文题目}}": "topic", "{{摘要}}": "abstract", "{{关键词}}": "keywords"}}
    state = {"topic": "预约系统设计", "workspace_dir": str(root), "new_template_path": str(template),
             "single_source_of_truth": {"tech_stack": "Python", "roles": ["学生", "管理员"]},
             "outline_plan": [{"section_id":"1", "title":"第一章 系统设计", "target_words":0, "status":"pending", "planned_assets":[]},
                              {"section_id":"1.1", "title":"预约流程", "target_words":100, "status":"pending", "planned_assets":[]}],
             "bib_pool": [], "assembly_contract": contract}
    return state, contract


def commit_fixture(state, assets=()):
    root = Path(state["workspace_dir"])
    content = "系统采用统一预约状态管理，学生提交预约后由服务端校验时段与角色权限。已占用的时段不能重复预约，取消操作需要验证预约归属。管理员负责查看预约信息并进行必要维护。该设计用于阐述业务规则，尚未进行性能测量。"
    for asset in assets:
        marker = "TABLE" if asset["kind"] == "table" else "FIG"
        content += f"如[[REF_{marker}:{asset['asset_id']}]]所示。"
    state["outline_plan"][-1]["planned_assets"] = [a["asset_id"]+" "+a["description"] for a in assets]
    record = {"section_id":"1.1", "title":"预约流程", "summary":"说明角色、时段及预约约束。",
              "key_facts":["学生和管理员"], "content":content, "citations":[], "assets":list(assets),
              "content_path":"writing/sections/1.1.md", "content_sha256":sha256(content.encode()).hexdigest(), "node_index":1}
    record["asset_sha256"] = {a[k]: sha256((root/a[k]).read_bytes()).hexdigest() for a in assets for k in ("path", "source_path")}
    path = root / record["content_path"]; path.parent.mkdir(parents=True, exist_ok=True); path.write_text(content)
    state["writing_records"] = [record]
    return state
