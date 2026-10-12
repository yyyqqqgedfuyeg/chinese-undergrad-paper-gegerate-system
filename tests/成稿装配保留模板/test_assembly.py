import copy
from hashlib import sha256
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from PIL import Image
from agent.assembly import AssemblyAgent
from agent.assembly.engine import package, parse, NS, text
from tests.support.paper_fixture import fixture, commit_fixture


class AssemblyTests(unittest.TestCase):
    def test_template_formats_content_and_determinism(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); state, contract = fixture(root)
            assets = root/"writing/assets"; assets.mkdir(parents=True)
            (assets/"table.json").write_text(json.dumps({"columns":["角色", "操作"],"rows":[["学生", "预约"],["管理员", "维护"]]}, ensure_ascii=False))
            Image.new("RGB", (640,320), "blue").save(assets/"figure.png")
            (assets/"source.py").write_text("# 测试图片源码")
            state = commit_fixture(state, [
                {"asset_id":"table_1", "description":"权限表", "kind":"table", "source_path":"writing/assets/table.json", "path":"writing/assets/table.json"},
                {"asset_id":"fig_1", "description":"业务流程", "kind":"diagram", "source_path":"writing/assets/source.py", "path":"writing/assets/figure.png"}])
            state["layout_options"] = {"chapter_page_break": True}
            output = root/"assembly/final.docx"
            agent = AssemblyAgent(state, contract, "上游提交的摘要", ["预约", "权限"], str(output))
            result = agent.invoke(); first = output.read_bytes()
            self.assertTrue(result["template_format_preserved"])
            self.assertTrue(result["deterministic_replay_identical"])
            agent.invoke(); self.assertEqual(first, output.read_bytes())
            original, final = package(state["new_template_path"]), package(output)
            for name in ("word/styles.xml", "word/theme/theme1.xml", "word/fontTable.xml", "word/numbering.xml"):
                self.assertEqual(original[name], final[name], name)
            self.assertEqual(sha256(Path(state["new_template_path"]).read_bytes()).hexdigest(), contract["template_sha256"])
            doc = Document(output)
            joined = text(parse(final["word/document.xml"]))
            self.assertIn("学校声明：不得删除", joined)
            self.assertIn("上游提交的摘要", joined)
            self.assertIn("如表1所示。如图1所示。", joined)
            self.assertNotIn("{{", joined)
            tree = parse(final["word/document.xml"])
            fields = tree.xpath(".//w:fldSimple/@w:instr", namespaces=NS)
            self.assertEqual(sum(x.startswith("REF PaperAsset") for x in fields), 2)
            self.assertEqual(sum(x.startswith("SEQ Paper") for x in fields), 2)
            self.assertEqual(len(tree.xpath(".//w:bookmarkStart", namespaces=NS)), 2)
            self.assertEqual(len(doc.inline_shapes), 1)
            figure_p = tree.xpath(".//w:p[.//w:drawing]", namespaces=NS)[0]
            spacing = figure_p.find("w:pPr/w:spacing", NS)
            self.assertEqual(spacing.get("{" + NS["w"] + "}lineRule"), "auto")
            self.assertEqual(spacing.get("{" + NS["w"] + "}line"), "240")
            self.assertTrue(tree.xpath(".//w:pageBreakBefore[@w:val='1']", namespaces=NS))

            self.assertEqual([[c.text for c in row.cells] for row in doc.tables[0].rows], [["角色","操作"],["学生","预约"],["管理员","维护"]])
            old = Document(state["new_template_path"])
            proto = next(p for p in old.paragraphs if p.text == "{{原型:body}}")
            generated = next(p for p in doc.paragraphs if p.text.startswith("系统采用"))
            from lxml import etree as E
            self.assertEqual(E.tostring(proto._p.pPr), E.tostring(generated._p.pPr))
            self.assertEqual(E.tostring(proto.runs[0]._r.rPr), E.tostring(generated.runs[0]._r.rPr))
            self.assertEqual(doc.sections[0].left_margin, old.sections[0].left_margin)
            self.assertEqual(doc.sections[0].header.paragraphs[0].text, "学校模板 预约系统设计")
            tools={t.name:t for t in agent.get_all_tools()}
            output.write_bytes(b"tampered")
            with self.assertRaises(ValueError): tools["verify_assembly"].invoke({})

    def test_missing_prototype_tampered_content_and_output_boundary(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); state, contract = fixture(root); commit_fixture(state)
            bad = copy.deepcopy(contract); del bad["paragraphs"]["body"]
            with self.assertRaises(ValueError):
                AssemblyAgent(state,bad,"摘要",["关键词"],str(root/"assembly/x.docx")).invoke()
            with self.assertRaises(PermissionError):
                AssemblyAgent(state,contract,"摘要",["关键词"],state["new_template_path"])
            (root/"writing/sections/1.1.md").write_text("被修改的正文")
            with self.assertRaises(ValueError):
                AssemblyAgent(state,contract,"摘要",["关键词"],str(root/"assembly/x.docx")).invoke()

    def test_template_hash_mismatch_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); state, contract = fixture(root); commit_fixture(state)
            contract["template_sha256"] = "stale"
            with self.assertRaises(ValueError):
                AssemblyAgent(state,contract,"摘要",["关键词"],str(root/"assembly/x.docx")).invoke()

    def test_repeated_citations_are_native_and_follow_first_use(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder); state,contract=fixture(root); commit_fixture(state)
            state["bib_pool"]=[{"key":"unused","title":"未引文献","year":2024,"authors":"甲"},
                               {"key":"b","title":"设计依据","year":2025,"authors":"乙","formatted":"乙. 设计依据[J]. 示例期刊, 2025."}]
            record=state["writing_records"][0]
            record["content"] += "依据[[REF_CITE:b]]，再次引用[[REF_CITE:b]]。"
            record["citations"]=["b"]
            record["content_sha256"]=sha256(record["content"].encode()).hexdigest()
            (root/record["content_path"]).write_text(record["content"])
            output=root/"assembly/final.docx"
            AssemblyAgent(state,contract,"摘要",["关键词"],str(output)).invoke()
            tree=parse(package(output)["word/document.xml"])
            self.assertIn("依据[1]，再次引用[1]。", text(tree))
            self.assertIn("[1] 乙. 设计依据",text(tree))
            self.assertNotIn("未引文献",text(tree))
            fields=tree.xpath(".//w:fldSimple/@w:instr",namespaces=NS)
            self.assertEqual(len(fields),2)
            self.assertEqual(fields[0],fields[1])
            for field in tree.xpath(".//w:fldSimple", namespaces=NS):
                self.assertEqual(field.find("w:r/w:rPr/w:vertAlign", NS).get("{"+NS["w"]+"}val"), "superscript")
            entry = next(p for p in tree.xpath(".//w:body/w:p", namespaces=NS) if text(p).startswith("[1] 乙"))
            self.assertEqual(entry.find("w:pPr/w:outlineLvl", NS).get("{"+NS["w"]+"}val"), "9")
            self.assertFalse(entry.xpath(".//w:vertAlign[@w:val='superscript']", namespaces=NS))


    def test_asset_mutation_rejected_after_commit(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder); state,contract=fixture(root)
            asset=root/"writing/assets/table.json";asset.parent.mkdir(parents=True)
            asset.write_text(json.dumps({"columns":["角色","操作"],"rows":[["学生","预约"]]}))
            commit_fixture(state,[{"asset_id":"tbl_1","kind":"table","description":"角色表",
                                  "path":"writing/assets/table.json","source_path":"writing/assets/table.json"}])
            asset.write_text(json.dumps({"columns":["角色","操作"],"rows":[["学生","越权操作"]]}))
            with self.assertRaisesRegex(ValueError,"资产被修改"):
                AssemblyAgent(state,contract,"摘要",["关键词"],str(root/"assembly/final.docx")).invoke()

    def test_unbound_school_template_empty_headings_are_rejected(self):
        from agent.paths import project_root
        from agent.assembly.engine import archive_bytes, xml, set_text, W
        from lxml import etree as E
        source=project_root()/"tests/模板提取/input/附件9 正文格式模板.docx"
        if not source.exists():
            self.skipTest("本地学校模板未提供")
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder); state,_=fixture(root); commit_fixture(state)
            original=package(source)
            parts=dict(original); doc=parse(parts["word/document.xml"])
            paragraphs=doc.find("w:body",NS).findall("w:p",NS)
            # 准备阶段仅标识原有段落，不修改学校格式；原文件始终保留。
            indices={"heading1":91,"heading2":93,"body":97}
            mapping={name:"{{学校原型:"+name+"}}" for name in indices}
            for name,index in indices.items():set_text(paragraphs[index],mapping[name])
            set_text(paragraphs[90],"{{正文}}")
            set_text(paragraphs[73],"{{论文题目}}")
            set_text(paragraphs[74],"摘  要：{{摘要}}")
            set_text(paragraphs[78],"关键词：{{关键词}}")
            parts["word/document.xml"]=xml(doc)
            template=root/"school-approved.docx";template.write_bytes(archive_bytes(parts))
            contract={"template_sha256":sha256(template.read_bytes()).hexdigest(),"body_marker":"{{正文}}",
                      "paragraphs":mapping,"fields":{"{{论文题目}}":"topic","{{摘要}}":"abstract","{{关键词}}":"keywords"}}
            state["new_template_path"]=str(template)
            output=root/"assembly/school.docx"
            # This partial contract leaves raw template headings unfilled.
            # The new pre-delivery gate must reject them instead of shipping a skeleton.
            with self.assertRaisesRegex(ValueError, "标题下没有有效内容"):
                AssemblyAgent(state,contract,"学校模板装配验证摘要",["模板","格式"],str(output)).invoke()
            self.assertFalse(output.exists())
            self.assertEqual(package(source), original)

    def test_duplicate_reference_prefix_is_rejected_before_assembly(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder); state,contract=fixture(root)
            asset=root/"writing/assets/table.json";asset.parent.mkdir(parents=True)
            asset.write_text(json.dumps({"columns":["角色","操作"],"rows":[["学生","预约"]]}))
            commit_fixture(state,[{"asset_id":"tbl_1","kind":"table","description":"角色表",
                                  "path":"writing/assets/table.json","source_path":"writing/assets/table.json"}])
            record=state["writing_records"][0]
            record["content"]=record["content"].replace('如[[REF_TABLE:', '如表[[REF_TABLE:')
            record["content_sha256"]=sha256(record["content"].encode()).hexdigest()
            (root/record["content_path"]).write_text(record["content"])
            with self.assertRaisesRegex(ValueError,"前缀"):
                AssemblyAgent(state,contract,"摘要",["关键词"],str(root/"assembly/final.docx")).invoke()
