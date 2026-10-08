"""Deterministic regression tests for template cleaning and format-preserving fill.

.venv/bin/python -m unittest test.test_template_pipeline -v
"""
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.shared import Pt
from lxml import etree as ET

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('template_pipeline',ROOT/'skills/docx/scripts/template_pipeline.py')
pipeline = importlib.util.module_from_spec(spec)
spec.loader.exec_module(pipeline)


class TemplatePipelineTests(unittest.TestCase):
    def test_fragmented_replacement_preserves_surrounding_runs(self):
        d = Document()
        p = d.add_paragraph()
        p.add_run('前缀 {{姓').bold = True
        p.add_run('名}} 后缀').italic = True
        pipeline.replace_span(p._p,'{{姓名}}','张三')
        self.assertEqual(p.text,'前缀 张三 后缀')
        self.assertTrue(p.runs[0].bold)
        self.assertTrue(p.runs[1].italic)
        with self.assertRaises(ValueError):
            pipeline.replace_span(p._p,'不存在','x')

    def test_fill_preserves_formats_and_requires_all_values(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            d = Document()
            p = d.add_paragraph()
            p.add_run('{{标').font.size = Pt(16)
            p.add_run('题}}').font.size = Pt(10)
            d.add_table(rows=1,cols=1).cell(0,0).text = '{{表格}}'
            d.sections[0].header.paragraphs[0].text = '论文题目：{{页眉题目}}'
            original = folder/'original.docx'
            d.save(original)
            values = folder/'values.json'
            values.write_text(json.dumps({'{{标题}}':'示例标题','{{表格}}':'示例表格','{{页眉题目}}':'示例标题'}))
            filled = folder/'filled.docx'
            pipeline.fill(original,values,filled)
            result = Document(filled)
            self.assertEqual(result.paragraphs[0].text,'示例标题')
            self.assertEqual(result.paragraphs[0].runs[0].font.size.pt,16)
            self.assertEqual(result.paragraphs[0].runs[1].font.size.pt,10)
            self.assertEqual(result.tables[0].cell(0,0).text,'示例表格')
            self.assertEqual(result.sections[0].header.paragraphs[0].text,'论文题目：示例标题')
            before,after = pipeline.package(original),pipeline.package(filled)
            for name in before:
                if name not in ('word/document.xml','word/header1.xml'):
                    self.assertEqual(before[name],after[name],name)
            values.write_text(json.dumps({'{{标题}}':'示例标题'}))
            with self.assertRaises(ValueError):
                pipeline.fill(original,values,filled)
            values.write_text(json.dumps({'{{标题}}':'示例标题','{{表格}}':'示例\n表格'}))
            with self.assertRaises(ValueError):
                pipeline.fill(original,values,filled)

    def test_plan_rejects_stale_text_and_duplicate_edits(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'test.docx'
            d = Document()
            d.add_paragraph('指导：***')
            d.save(path)
            parts = pipeline.package(path)
            edit = {'paragraph':0,'expected_text':'旧原文','evidence':'测试','replacements':[]}
            with self.assertRaises(ValueError):
                pipeline.apply_plan(parts,{'edits':[edit]})
            edit['expected_text'] = '指导：***'
            with self.assertRaises(ValueError):
                pipeline.apply_plan(parts,{'edits':[edit,edit]})

    def test_grid_override_is_inserted_before_spacing_and_run_defaults(self):
        d = Document()
        p = d.add_paragraph('test')
        p.paragraph_format.line_spacing = 1.5
        p._p.get_or_add_pPr().append(ET.Element(pipeline.q('rPr')))
        pipeline.set_props(p._p,'pPr',{'snapToGrid':{'val':'0'}})
        tags = [ET.QName(c).localname for c in p._p.pPr]
        self.assertLess(tags.index('snapToGrid'),tags.index('spacing'))
        self.assertLess(tags.index('snapToGrid'),tags.index('rPr'))

    def test_style_order_repair_preserves_property_values(self):
        data = ('<w:styles xmlns:w="'+pipeline.W+'"><w:style w:styleId="1">'
                '<w:name w:val="标题"/><w:semiHidden/><w:uiPriority w:val="9"/>'
                '<w:rPr><w:sz w:val="24"/></w:rPr></w:style></w:styles>').encode()
        repaired,ids = pipeline.normalize_style_order(data)
        self.assertEqual(ids,['1'])
        style = pipeline.parse(repaired)[0]
        self.assertEqual([ET.QName(c).localname for c in style],['name','uiPriority','semiHidden','rPr'])
        self.assertEqual(style.find('w:rPr/w:sz',pipeline.NS).get(pipeline.q('val')),'24')
        self.assertEqual(pipeline.normalize_style_order(repaired),(repaired,[]))

    def test_rebuilt_styles_are_visible_applied_and_editable(self):
        source=ROOT/'test/tmp/附件9 正文格式模板.docx'
        if not source.exists():self.skipTest('Annex 9 source unavailable')
        plan=pipeline.load(ROOT/'test/fixtures/annex9-cleaning-plan.json')
        plan['style_rebuild']=pipeline.load(ROOT/'test/fixtures/annex9-style-rebuild.json')
        parts,_,report=pipeline.transform(pipeline.package(source),plan)
        self.assertEqual(report['removed_style_count'],67)
        definitions=pipeline.parse(parts['word/styles.xml'])
        self.assertFalse(set(report['removed_style_ids']) & {s.get(pipeline.q('styleId')) for s in definitions.findall('w:style',pipeline.NS)})
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'named.docx';pipeline.write_package(path,parts)
            d=Document(path)
            self.assertEqual(d.paragraphs[88].style.name,'论文正文')
            self.assertEqual(d.paragraphs[91].style.name,'论文一级标题')
            body=d.styles['论文正文']
            self.assertTrue(body.quick_style)
            r=next(r for r in d.paragraphs[88].runs if '{{' in r.text)
            self.assertIsNone(r.font.size)
            self.assertEqual(body.font.size.pt,12)
            body.font.size=Pt(14);d.save(path)
            reloaded=Document(path)
            self.assertEqual(reloaded.paragraphs[88].style.font.size.pt,14)
            self.assertIsNone(next(r for r in reloaded.paragraphs[88].runs if '{{' in r.text).font.size)
        self.assertTrue(all(s['applied_paragraph_count']>0 for s in report['new_paragraph_styles']))
        errors=[];pipeline.verify_named_styles(parts,plan,errors)
        self.assertEqual(errors,[])

    def test_verifier_detects_media_tampering(self):
        source = ROOT/'test/tmp/附件9 正文格式模板.docx'
        plan = ROOT/'test/fixtures/annex9-cleaning-plan.json'
        if not source.exists() or not plan.exists():
            self.skipTest('Annex 9 source or accepted real-agent plan unavailable')
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder)
            pipeline.build(source,out,plan)
            self.assertTrue(pipeline.verify(source,out,plan)['passed'])
            parts = pipeline.package(out/'clean-template.docx')
            media = next(n for n in parts if n.startswith('word/media/'))
            parts[media] += b'tampered'
            pipeline.write_package(out/'clean-template.docx',parts)
            with self.assertRaises(SystemExit):
                pipeline.verify(source,out,plan)


if __name__ == '__main__':
    unittest.main()
