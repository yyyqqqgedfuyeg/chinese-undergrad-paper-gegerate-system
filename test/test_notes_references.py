"""Regression tests for native footnotes and cross-reference construction."""
from copy import deepcopy
import hashlib
import importlib.util
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('nr',ROOT/'skills/docx/scripts/notes_references.py')
nr=importlib.util.module_from_spec(spec);spec.loader.exec_module(nr)


class NotesTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.base=Path(self.tmp.name);self.src=self.base/'source.docx';self.out=self.base/'output.docx'
        doc=Document()
        for sid in ('PaperBody','CaptionTest','ReferenceTest'):
            doc.styles.add_style(sid,WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph('preserved paragraph');doc.save(self.src)
        self.plan={'source_sha256':hashlib.sha256(self.src.read_bytes()).hexdigest(),
          'footnote_style':{'east_asia_font':'宋体','latin_font':'Times New Roman','size_pt':9,'marker_size_pt':9,'line':240},
          'blocks':[
            {'type':'paragraph','inlines':['Forward ',{'ref':'Table_A'},' note ',{'footnote':{'bookmark':'Note_A','text':'Native note content'}},' again ',{'ref':'Note_A','superscript':True},' citation ',{'ref':'Bib_A','superscript':True}]},
            {'type':'caption','style':'CaptionTest','bookmark':'Table_A','sequence':'Table','prefix':'表','text':'Results'},
            {'type':'table','rows':[['A','B'],['1','2']]},
            {'type':'caption','style':'ReferenceTest','bookmark':'Bib_A','sequence':'Bibliography','prefix':'[','suffix':']','text':'Demonstration only'}]}

    def test_native_notes_and_forward_repeated_references(self):
        original=self.src.read_bytes();report=nr.apply(self.src,self.plan,self.out)
        self.assertTrue(report['passed']);self.assertEqual(report['footnote_count'],1)
        self.assertEqual([r['kind'] for r in report['references']],['REF','NOTEREF','REF'])
        self.assertEqual(original,self.src.read_bytes())
        parts=nr.read(self.out);foot=nr.parse(parts['word/footnotes.xml'])
        self.assertEqual(len(foot.findall('w:footnote',nr.NS)),3)
        self.assertIn(b'footnotes.xml',parts['word/_rels/document.xml.rels'])
        self.assertIn(b'footnotes+xml',parts['[Content_Types].xml'])
        self.assertEqual(len(nr.parse(parts['word/document.xml']).findall('.//w:tbl',nr.NS)),1)

    def test_note_style_is_visible_and_editable(self):
        nr.apply(self.src,self.plan,self.out);doc=Document(self.out)
        style=doc.styles['论文脚注正文'];self.assertEqual(style.font.size.pt,9);style.font.size=Pt(11)
        self.assertTrue(style.quick_style)
        foot=nr.parse(nr.read(self.out)['word/footnotes.xml'])
        text=next(r for r in foot.findall('.//w:r',nr.NS) if 'Native' in ''.join(r.itertext()))
        self.assertIsNone(text.find('w:rPr/w:sz',nr.NS))
        self.assertEqual(foot.find('.//w:pPr/w:pStyle',nr.NS).get(nr.q('val')),'PaperFootnoteText')

    def test_duplicate_reference_prefix_fails(self):
        self.plan['blocks'][0]['inlines'][0]='See 表'
        with self.assertRaisesRegex(ValueError,'already includes prefix'):nr.apply(self.src,self.plan,self.out)

    def test_pdf_adapter_retains_authored_fields_and_styles(self):
        nr.apply(self.src,self.plan,self.out);original=self.out.read_bytes()
        render=self.base/'pdf-input.docx';report=nr.pdf_input(self.out,render)
        self.assertEqual(self.out.read_bytes(),original)
        self.assertEqual(len(report['converted_note_references']),1)
        self.assertIn('NOTEREF', ' '.join(nr.inspect(self.out)['fields']))
        self.assertNotIn('NOTEREF', ' '.join(nr.inspect(render)['fields']))
        self.assertTrue(nr.inspect(render)['passed'])
        self.assertEqual(nr.read(render)['word/styles.xml'],nr.read(self.out)['word/styles.xml'])
        d=nr.parse(nr.read(render)['word/document.xml'])
        link=d.find('.//w:hyperlink',nr.NS)
        self.assertEqual(link.get(nr.q('anchor')),'Note_A')
        self.assertEqual(''.join(link.itertext()),'1')

    def test_stale_hash(self):
        self.plan['source_sha256']='wrong'
        with self.assertRaisesRegex(ValueError,'source_sha256'):nr.apply(self.src,self.plan,self.out)
        self.assertFalse(self.out.exists())

    def test_duplicate_bookmark(self):
        self.plan['blocks'][-1]['bookmark']='Table_A'
        with self.assertRaisesRegex(ValueError,'duplicate bookmark'):nr.apply(self.src,self.plan,self.out)

    def test_missing_reference_fails(self):
        self.plan['blocks'][0]['inlines'][1]['ref']='Missing'
        with self.assertRaisesRegex(ValueError,'Unknown reference'):nr.apply(self.src,self.plan,self.out)

    def test_inspection_rejects_broken_note_reference(self):
        nr.apply(self.src,self.plan,self.out);parts=nr.read(self.out)
        root=nr.parse(parts['word/document.xml']);root.find('.//w:footnoteReference',nr.NS).set(nr.q('id'),'999')
        parts['word/document.xml']=nr.dump(root)
        with ZipFile(self.out,'w') as z:
            for n,data in parts.items():z.writestr(n,data)
        self.assertFalse(nr.inspect(self.out)['passed'])

    def test_existing_sequence_not_silently_misnumbered(self):
        nr.apply(self.src,self.plan,self.out)
        parts=nr.read(self.out)
        # Reuse output as source after removing the new style definitions only.
        st=nr.parse(parts['word/styles.xml'])
        for s in list(st):
            if s.get(nr.q('styleId')) in ('PaperFootnoteText','PaperFootnoteReference'):st.remove(s)
        parts['word/styles.xml']=nr.dump(st)
        modified=self.base/'existing.docx'
        with ZipFile(modified,'w') as z:
            for n,data in parts.items():z.writestr(n,data)
        self.plan['source_sha256']=hashlib.sha256(modified.read_bytes()).hexdigest()
        for b in self.plan['blocks']:
            if 'bookmark' in b:b['bookmark']+='_New'
            for item in b.get('inlines',[]):
                if isinstance(item,dict) and 'footnote' in item:item['footnote']['bookmark']+='_New'
        with self.assertRaisesRegex(ValueError,'Existing sequence'):nr.apply(modified,self.plan,self.base/'second.docx')

    def test_refresh_uses_note_order_not_internal_id_and_preserves_parts(self):
        nr.apply(self.src,self.plan,self.out);parts=nr.read(self.out);d=nr.parse(parts['word/document.xml'])
        foot=nr.parse(parts['word/footnotes.xml']);body=d.find('w:body',nr.NS)
        p=nr.node('p');r=nr.run(style='PaperFootnoteReference');r.append(nr.node('footnoteReference',id=9));p.append(r);body.insert(1,p)
        fn=deepcopy(next(n for n in foot if n.get(nr.q('id'))=='1'));fn.set(nr.q('id'),'9');foot.append(fn)
        for name in ('Table_A','Bib_A'):
            start=next(n for n in d.findall('.//w:bookmarkStart',nr.NS) if n.get(nr.q('name'))==name)
            p=start.getparent();copy=deepcopy(p);b=copy.find('w:bookmarkStart',nr.NS);b.set(nr.q('name'),name+'_Earlier')
            bid='90' if name=='Table_A' else '91';b.set(nr.q('id'),bid);copy.find('w:bookmarkEnd',nr.NS).set(nr.q('id'),bid);body.insert(body.index(p),copy)
        parts['word/document.xml']=nr.dump(d);parts['word/footnotes.xml']=nr.dump(foot)
        mutated=self.base/'mutated.docx'
        with ZipFile(mutated,'w') as z:
            for name,data in parts.items():z.writestr(name,data)
        refreshed=self.base/'refreshed.docx';report=nr.refresh(mutated,refreshed)
        self.assertTrue(report['passed'])
        changes={c['instruction']:c['after'] for c in report['cache_updates']}
        self.assertEqual(changes['NOTEREF Note_A \\h \\f'],'2')
        self.assertEqual(changes['REF Table_A \\h'],'表2')
        self.assertEqual(changes['REF Bib_A \\h'],'[2]')
        saved=nr.read(refreshed)
        for name,data in parts.items():
            if name!='word/document.xml':self.assertEqual(data,saved[name])
        self.assertEqual(nr.refresh(refreshed,self.base/'again.docx')['cache_updates'],[])

    def test_unsupported_sequence_switches_fail(self):
        nr.apply(self.src,self.plan,self.out);parts=nr.read(self.out);d=nr.parse(parts['word/document.xml'])
        f=next(f for f in d.findall('.//w:fldSimple',nr.NS) if 'SEQ Table' in f.get(nr.q('instr')))
        f.set(nr.q('instr'),' SEQ Table \\r 9 ');parts['word/document.xml']=nr.dump(d)
        with ZipFile(self.out,'w') as z:
            for name,data in parts.items():z.writestr(name,data)
        with self.assertRaisesRegex(ValueError,'Unsupported SEQ'):nr.refresh(self.out,self.base/'bad.docx')

    def test_ambiguous_anchor(self):
        self.plan['insert_before']='missing anchor'
        with self.assertRaisesRegex(ValueError,'exactly one'):nr.apply(self.src,self.plan,self.out)


if __name__=='__main__':unittest.main()
